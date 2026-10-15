"""يولّد Jisr/AwamilJadwal.lean من مودَع «جدول العوامل المئة» (xlsx المالك؛ بصمته في mudaat/sijill.json).

التعيين منقول بنصّه من مولّد المالك (slge-closed-v13/awamil_xlsx/gen_awamil.py): يُستعمل عمودان فقط — العامل، والحكم
الإعرابيّ (والأسلوب للاستفهام والنداء) — والصنف رقمٌ من 13 بحسب الحكم الإعرابيّ، و«مضمرة» في الحكم تعلّم العامل المضمَر.
الأعمدةُ الأخرى (الأوزان، CV، الباب، الشواهد، FOL) مستبعدة: ليست من حكم العامل.
التشغيل: python3 -I scripts/bina_awamil.py <awamil_mia.xlsx> [<Jisr/AwamilJadwal.lean> [<Jisr/AwamilUslub.lean>]]
"""
import hashlib
import os
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def cls(r):
    name, style, h = r[0], r[1], r[2]
    if style == "الاستفهام": return 12
    if style == "النداء": return 13
    if h == "تنصب الاسم وترفع الخبر": return 2
    if h == "ترفع الاسم وتنصب الخبر": return 1
    if h == "تنصب مفعولين": return 3
    if h.startswith("ترفع الاسم وترفع/تنصب الخبر"): return 6
    if h == "تجر الاسم وتنصب الفعل المضارع": return 9
    if h.startswith("تجرّ") or h == "تدخل على لفظ الجلالة": return 7
    if h.startswith("تعطف"): return 8
    if h.startswith("تجزم المضارع"): return 5
    if h.startswith("تنصب المضارع"): return 4
    if h == "أداة شرط جزم فعلين": return 10
    if h.startswith("ظرف/حرف شرط غير جازم"): return 11
    return 0


def wallid(path):
    sha = hashlib.sha256(open(path, "rb").read()).hexdigest()
    rows = [[("" if v is None else str(v)) for v in r]
            for r in openpyxl.load_workbook(path, data_only=True).active.iter_rows(values_only=True)][1:]
    tab = [(r[0], cls(r), "مضمرة" in r[2]) for r in rows]
    esc = lambda x: x.replace('"', '\\"')
    body = ",\n  ".join(f'("{esc(n)}", {c}, {"true" if h else "false"})' for n, c, h in tab)
    text = f'''/-! جدولٌ مولَّد بـ scripts/bina_awamil.py من مودَع «جدول العوامل المئة» (sha256 {sha}) — لا يُحرَّر يدويًّا.
التعيينُ تعيينُ المالك (awamil_xlsx/gen_awamil.py في حزمته slge-closed-v13). (العامل، الصنف بالحكم الإعرابيّ، العاملُ مضمَر؟). -/
namespace Jisr
namespace Awamil

def table : List (String × Nat × Bool) := [
  {body}
]

end Awamil
end Jisr
'''
    return sha, tab, text


def wallidUslub(path):
    """جدول ثانٍ: (العامل، الأسلوب النحوي) من عمود «الأسلوب النحوي» بلفظه — لاستنفاد أساليب الإخبار والإنشاء."""
    sha = hashlib.sha256(open(path, "rb").read()).hexdigest()
    rows = [[("" if v is None else str(v)) for v in r]
            for r in openpyxl.load_workbook(path, data_only=True).active.iter_rows(values_only=True)][1:]
    esc = lambda x: x.replace('"', '\\"')
    body = ",\n  ".join(f'("{esc(r[0])}", "{esc(r[1])}")' for r in rows)
    return f'''/-! جدولٌ مولَّد بـ scripts/bina_awamil.py من مودَع «جدول العوامل المئة» (sha256 {sha}) — لا يُحرَّر يدويًّا.
(العامل، الأسلوب النحويّ) بلفظ العمود. -/
namespace Jisr
namespace Awamil

def uslub : List (String × String) := [
  {body}
]

end Awamil
end Jisr
'''


def main():
    path = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) > 2 else os.path.join(ROOT, "Jisr", "AwamilJadwal.lean")
    sha, tab, text = wallid(path)
    open(out, "w", encoding="utf-8").write(text)
    out2 = sys.argv[3] if len(sys.argv) > 3 else os.path.join(ROOT, "Jisr", "AwamilUslub.lean")
    open(out2, "w", encoding="utf-8").write(wallidUslub(path))
    from collections import Counter
    print(sha[:16], len(tab), sorted(Counter(c for _, c, _ in tab).items()))


if __name__ == "__main__":
    main()
